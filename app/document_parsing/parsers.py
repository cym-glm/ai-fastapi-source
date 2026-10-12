import hashlib
from abc import ABC, abstractmethod
from pathlib import Path
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
import fitz

from app.document_parsing.cleaner import TextCleaner
from app.document_parsing.schemas import (
    ParsedDocument,
    ParsedElement,
    ParsedElementType,
    ParseStatus,
)


def build_document_id(
    file_path: str | Path,
) -> str:
    path = Path(file_path)

    raw = f"{path.name}:{path.stat().st_size}:{path.stat().st_mtime}"
    return hashlib.md5(raw.encode("utf-8")).hexdigest()


class BaseDocumentParser(ABC):
    def __init__(self):
        self.cleaner = TextCleaner()

    @abstractmethod
    def parse(
        self,
        file_path: str | Path,
        tenant_id: str = "tenant_001",
        created_by: str = "system",
    ) -> ParsedDocument:
        raise NotImplementedError

    def build_base_metadata(
        self,
        file_path: str | Path,
        tenant_id: str,
        created_by: str,
    ) -> dict:
        path = Path(file_path)

        return {
            "tenant_id": tenant_id,
            "created_by": created_by,
            "file_name": path.name,
            "file_size": path.stat().st_size,
            "source_path": str(path),
        }


class TxtDocumentParser(BaseDocumentParser):
    def parse(
        self,
        file_path: str | Path,
        tenant_id: str = "tenant_001",
        created_by: str = "system",
    ) -> ParsedDocument:
        path = Path(file_path)

        try:
            text = path.read_text(
                encoding="utf-8"
            )
        except UnicodeDecodeError:
            text = path.read_text(
                encoding="gbk",
                errors="ignore",
            )

        text = self.cleaner.clean_text(text)

        element = ParsedElement(
            element_id=f"{path.stem}_txt_0",
            type=ParsedElementType.PARAGRAPH,
            content=text,
            order_index=0,
            metadata={
                "parser": "TxtDocumentParser",
            },
        )

        return ParsedDocument(
            document_id=build_document_id(path),
            title=path.stem,
            source_path=str(path),
            file_name=path.name,
            file_type=path.suffix.lower(),
            text=text,
            elements=[element] if text else [],
            status=ParseStatus.SUCCESS if text else ParseStatus.FAILED,
            metadata=self.build_base_metadata(
                path,
                tenant_id,
                created_by,
            ),
            errors=[] if text else ["文本内容为空"],
        )


class MarkdownDocumentParser(BaseDocumentParser):
    def parse(
        self,
        file_path: str | Path,
        tenant_id: str = "tenant_001",
        created_by: str = "system",
    ) -> ParsedDocument:
        path = Path(file_path)
        raw_text = path.read_text(
            encoding="utf-8",
            errors="ignore",
        )

        elements: list[ParsedElement] = []
        current_code_lines: list[str] = []
        in_code_block = False

        lines = raw_text.splitlines()

        for index, line in enumerate(lines):
            stripped = line.strip()

            if stripped.startswith("```"):
                if not in_code_block:
                    in_code_block = True
                    current_code_lines = []
                else:
                    in_code_block = False
                    content = "\n".join(current_code_lines).strip()

                    if content:
                        elements.append(
                            ParsedElement(
                                element_id=f"{path.stem}_code_{index}",
                                type=ParsedElementType.CODE,
                                content=content,
                                order_index=len(elements),
                                metadata={
                                    "parser": "MarkdownDocumentParser",
                                },
                            )
                        )

                continue

            if in_code_block:
                current_code_lines.append(line)
                continue

            if not stripped:
                continue

            if stripped.startswith("#"):
                content = stripped.lstrip("#").strip()
                element_type = ParsedElementType.TITLE
            elif stripped.startswith(("-", "*", "+")):
                content = stripped
                element_type = ParsedElementType.LIST
            else:
                content = stripped
                element_type = ParsedElementType.PARAGRAPH

            elements.append(
                ParsedElement(
                    element_id=f"{path.stem}_md_{index}",
                    type=element_type,
                    content=content,
                    order_index=len(elements),
                    metadata={
                        "parser": "MarkdownDocumentParser",
                    },
                )
            )

        text = self.cleaner.clean_text(
            "\n\n".join(
                element.content
                for element in elements
            )
        )

        title = path.stem

        for element in elements:
            if element.type == ParsedElementType.TITLE:
                title = element.content
                break

        return ParsedDocument(
            document_id=build_document_id(path),
            title=title,
            source_path=str(path),
            file_name=path.name,
            file_type=path.suffix.lower(),
            text=text,
            elements=elements,
            status=ParseStatus.SUCCESS if elements else ParseStatus.FAILED,
            metadata=self.build_base_metadata(
                path,
                tenant_id,
                created_by,
            ),
            errors=[] if elements else ["Markdown 内容为空"],
        )

class HtmlDocumentParser(BaseDocumentParser):
    def parse(
        self,
        file_path: str | Path,
        tenant_id: str = "tenant_001",
        created_by: str = "system",
    ) -> ParsedDocument:
        path = Path(file_path)

        html = path.read_text(
            encoding="utf-8",
            errors="ignore",
        )

        soup = BeautifulSoup(
            html,
            "lxml",
        )

        for tag in soup([
            "script",
            "style",
            "noscript",
            "svg",
        ]):
            tag.decompose()

        title = path.stem

        if soup.title and soup.title.string:
            title = soup.title.string.strip()

        elements: list[ParsedElement] = []

        for index, tag in enumerate(
            soup.find_all(["h1", "h2", "h3", "p", "li", "td", "th"])
        ):
            content = self.cleaner.clean_text(
                tag.get_text(" ", strip=True)
            )

            if not content:
                continue

            if tag.name in ["h1", "h2", "h3"]:
                element_type = ParsedElementType.TITLE
            elif tag.name == "li":
                element_type = ParsedElementType.LIST
            elif tag.name in ["td", "th"]:
                element_type = ParsedElementType.TABLE
            else:
                element_type = ParsedElementType.PARAGRAPH

            elements.append(
                ParsedElement(
                    element_id=f"{path.stem}_html_{index}",
                    type=element_type,
                    content=content,
                    order_index=len(elements),
                    metadata={
                        "parser": "HtmlDocumentParser",
                        "tag": tag.name,
                    },
                )
            )

        text = self.cleaner.clean_text(
            "\n\n".join(
                element.content
                for element in elements
            )
        )

        return ParsedDocument(
            document_id=build_document_id(path),
            title=title,
            source_path=str(path),
            file_name=path.name,
            file_type=path.suffix.lower(),
            text=text,
            elements=elements,
            status=ParseStatus.SUCCESS if elements else ParseStatus.FAILED,
            metadata={
                **self.build_base_metadata(
                    path,
                    tenant_id,
                    created_by,
                ),
                "html_title": title,
            },
            errors=[] if elements else ["HTML 未解析出有效正文"],
        )


class DocxDocumentParser(BaseDocumentParser):
    def parse(
        self,
        file_path: str | Path,
        tenant_id: str = "tenant_001",
        created_by: str = "system",
    ) -> ParsedDocument:
        path = Path(file_path)
        doc = DocxDocument(str(path))

        elements: list[ParsedElement] = []

        for index, paragraph in enumerate(doc.paragraphs):
            content = self.cleaner.clean_text(
                paragraph.text
            )

            if not content:
                continue

            style_name = (
                paragraph.style.name
                if paragraph.style
                else ""
            )

            if style_name.startswith("Heading"):
                element_type = ParsedElementType.TITLE
            else:
                element_type = ParsedElementType.PARAGRAPH

            elements.append(
                ParsedElement(
                    element_id=f"{path.stem}_docx_p_{index}",
                    type=element_type,
                    content=content,
                    order_index=len(elements),
                    metadata={
                        "parser": "DocxDocumentParser",
                        "style": style_name,
                    },
                )
            )

        for table_index, table in enumerate(doc.tables):
            rows: list[str] = []

            for row in table.rows:
                cells = [
                    self.cleaner.clean_text(cell.text)
                    for cell in row.cells
                ]

                rows.append(" | ".join(cells))

            table_text = "\n".join(rows).strip()

            if table_text:
                elements.append(
                    ParsedElement(
                        element_id=f"{path.stem}_docx_table_{table_index}",
                        type=ParsedElementType.TABLE,
                        content=table_text,
                        order_index=len(elements),
                        metadata={
                            "parser": "DocxDocumentParser",
                            "table_index": table_index,
                        },
                    )
                )

        text = self.cleaner.clean_text(
            "\n\n".join(
                element.content
                for element in elements
            )
        )

        title = path.stem

        for element in elements:
            if element.type == ParsedElementType.TITLE:
                title = element.content
                break

        return ParsedDocument(
            document_id=build_document_id(path),
            title=title,
            source_path=str(path),
            file_name=path.name,
            file_type=path.suffix.lower(),
            text=text,
            elements=elements,
            status=ParseStatus.SUCCESS if elements else ParseStatus.FAILED,
            metadata=self.build_base_metadata(
                path,
                tenant_id,
                created_by,
            ),
            errors=[] if elements else ["DOCX 未解析出有效正文"],
        )



class PdfDocumentParser(BaseDocumentParser):
    def parse(
        self,
        file_path: str | Path,
        tenant_id: str = "tenant_001",
        created_by: str = "system",
    ) -> ParsedDocument:
        path = Path(file_path)
        elements: list[ParsedElement] = []
        errors: list[str] = []

        try:
            pdf = fitz.open(str(path))
        except Exception as exc:
            return ParsedDocument(
                document_id=build_document_id(path),
                title=path.stem,
                source_path=str(path),
                file_name=path.name,
                file_type=path.suffix.lower(),
                text="",
                elements=[],
                status=ParseStatus.FAILED,
                metadata=self.build_base_metadata(
                    path,
                    tenant_id,
                    created_by,
                ),
                errors=[f"PDF 打开失败：{exc}"],
            )

        for page_index in range(len(pdf)):
            page = pdf[page_index]

            try:
                page_text = page.get_text("text")
            except Exception as exc:
                errors.append(
                    f"第 {page_index + 1} 页文本提取失败：{exc}"
                )
                continue

            page_text = self.cleaner.clean_text(page_text)

            if not page_text:
                errors.append(
                    f"第 {page_index + 1} 页未提取到文本，可能是扫描件或图片页"
                )
                continue

            elements.append(
                ParsedElement(
                    element_id=f"{path.stem}_pdf_page_{page_index + 1}",
                    type=ParsedElementType.PAGE_TEXT,
                    content=page_text,
                    page_number=page_index + 1,
                    order_index=len(elements),
                    metadata={
                        "parser": "PdfDocumentParser",
                        "page_number": page_index + 1,
                    },
                )
            )

        text = self.cleaner.clean_text(
            "\n\n".join(
                element.content
                for element in elements
            )
        )

        metadata = {
            **self.build_base_metadata(
                path,
                tenant_id,
                created_by,
            ),
            "page_count": len(pdf),
        }

        pdf.close()

        if elements and errors:
            status = ParseStatus.PARTIAL_SUCCESS
        elif elements:
            status = ParseStatus.SUCCESS
        elif errors:
            status = ParseStatus.NEED_OCR
        else:
            status = ParseStatus.FAILED

        return ParsedDocument(
            document_id=build_document_id(path),
            title=path.stem,
            source_path=str(path),
            file_name=path.name,
            file_type=path.suffix.lower(),
            text=text,
            elements=elements,
            status=status,
            metadata=metadata,
            errors=errors,
        )