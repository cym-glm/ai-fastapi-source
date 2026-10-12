from pathlib import Path

from docx import Document

from app.document_parsing.service import document_parsing_service


def create_sample_docx():
    output_path = Path("data/sample_docs/contract_template.docx")
    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    doc = Document()

    doc.add_heading(
        "合同模板",
        level=1,
    )

    doc.add_paragraph(
        "本合同模板用于企业采购场景。"
    )

    doc.add_heading(
        "付款条款",
        level=2,
    )

    doc.add_paragraph(
        "甲方应在收到发票后 30 个自然日内完成付款。"
    )

    table = doc.add_table(
        rows=2,
        cols=3,
    )

    table.cell(0, 0).text = "条款"
    table.cell(0, 1).text = "说明"
    table.cell(0, 2).text = "备注"

    table.cell(1, 0).text = "付款周期"
    table.cell(1, 1).text = "30 个自然日"
    table.cell(1, 2).text = "以合同约定为准"

    doc.save(output_path)

    return output_path


def main():
    file_path = create_sample_docx()

    result = document_parsing_service.parse_file(
        file_path,
        tenant_id="tenant_001",
        created_by="u_10001",
    )

    print("status:", result.status)
    print("title:", result.title)

    print("=" * 80)
    print("elements:")
    for element in result.elements:
        print(
            element.type,
            element.metadata,
            "=>",
            element.content,
        )

    print("=" * 80)
    print("text:")
    print(result.text)


if __name__ == "__main__":
    main()